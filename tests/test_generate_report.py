import logging
import runpy
import sys
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from generate_report import (
    ReportError,
    build_report,
    build_summaries,
    clean_data,
    clean_dates,
    clean_money,
    clean_products,
    clean_regions,
    detect_separator,
    fill_blanks,
    find_missing,
    load_data,
    main,
    remove_duplicates,
)

SAMPLE_FILE = Path(__file__).parent.parent / "sample_csv_data.csv"

HEADER = "order_id,order_date,region,product,quantity,unit_price,revenue,sales_rep"

ROWS = [
    "ORD-1,2025-01-15,WEST,Laptop Stand,2,$45.00,90.00,Ana Lee",
    "ORD-2,02/10/2025, south ,laptop  stand,1,45.00,,",
    "ORD-3,05-Mar-2025,Eas,USB-C Hub,,39.50,79.00,Ben Cho",
    "ORD-3,05-Mar-2025,Eas,USB-C Hub,,39.50,79.00,Ben Cho",
]


def write_csv(folder: Path, separator: str = ",", name: str = "sales.csv") -> Path:
    path = folder / name
    lines = [HEADER, *ROWS]
    path.write_text("\n".join(line.replace(",", separator) for line in lines) + "\n")
    return path


def text_frame(**columns: list[str | None]) -> pd.DataFrame:
    return pd.DataFrame(columns, dtype="object")


@pytest.mark.parametrize("separator", [",", "\t", ";", "|"])
def test_detect_separator(tmp_path: Path, separator: str) -> None:
    path = tmp_path / "sales.txt"
    path.write_text(HEADER.replace(",", separator) + "\n")

    assert detect_separator(path) == separator


def test_detect_separator_ignores_commas_in_data(tmp_path: Path) -> None:
    path = tmp_path / "sales.txt"
    path.write_text(HEADER.replace(",", ";") + "\n" + "1;$2,790.00;$1,000.00\n")

    assert detect_separator(path) == ";"


@pytest.mark.parametrize("separator", [",", "\t", ";", "|"])
def test_load_data_reads_each_separator(tmp_path: Path, separator: str) -> None:
    df = load_data(write_csv(tmp_path, separator))

    assert len(df) == 4
    assert df.loc[0, "unit_price"] == "$45.00"


def test_load_data_reads_excel(tmp_path: Path) -> None:
    path = tmp_path / "sales.xlsx"
    pd.read_csv(write_csv(tmp_path), dtype=str).to_excel(path, index=False)

    df = load_data(path)

    assert len(df) == 4
    assert df.loc[0, "region"] == "WEST"


def test_load_data_cleans_column_names(tmp_path: Path) -> None:
    path = tmp_path / "sales.csv"
    header = (
        " Order_ID , ORDER_DATE ,Region,Product,Quantity,Unit_Price,Revenue,Sales_Rep"
    )
    path.write_text(header + "\n" + ROWS[0] + "\n")

    df = load_data(path)

    assert list(df.columns) == HEADER.split(",")


def test_load_data_missing_file(tmp_path: Path) -> None:
    with pytest.raises(ReportError, match="File not found"):
        load_data(tmp_path / "missing.csv")


def test_load_data_unsupported_type(tmp_path: Path) -> None:
    path = tmp_path / "sales.pdf"
    path.write_text("not a sales file")

    with pytest.raises(ReportError, match="Unsupported file type '.pdf'"):
        load_data(path)


def test_load_data_unreadable_file(tmp_path: Path) -> None:
    path = tmp_path / "sales.xlsx"
    path.write_text("not really an Excel file")

    with pytest.raises(ReportError, match="Could not read the file"):
        load_data(path)


def test_load_data_missing_columns(tmp_path: Path) -> None:
    path = tmp_path / "sales.csv"
    path.write_text("order_date,region,product\n2025-01-15,West,Laptop Stand\n")

    with pytest.raises(ReportError, match="quantity, unit_price, revenue, sales_rep"):
        load_data(path)


def test_load_data_no_rows(tmp_path: Path) -> None:
    path = tmp_path / "sales.csv"
    path.write_text(HEADER + "\n")

    with pytest.raises(ReportError, match="no data rows"):
        load_data(path)


def test_clean_regions() -> None:
    df = text_frame(region=["WEST", " west ", "Wes", "nor", "South", "EAST", "Middle"])

    result = clean_regions(df)["region"].tolist()

    assert result[:6] == ["West", "West", "West", "North", "South", "East"]
    assert pd.isna(result[6])


def test_clean_money() -> None:
    df = text_frame(
        unit_price=["$2,790.00", " 315.00 ", None],
        revenue=["$1,000.50", "abc", "20"],
    )

    result = clean_money(df)

    assert result["unit_price"].tolist()[:2] == [2790.0, 315.0]
    assert pd.isna(result.loc[2, "unit_price"])
    assert result.loc[0, "revenue"] == 1000.5
    assert pd.isna(result.loc[1, "revenue"])


def test_clean_dates() -> None:
    df = text_frame(
        order_date=["2025-01-15", "02/10/2025", "05-Mar-2025", "April 7, 2025", "soon"]
    )

    result = clean_dates(df)["order_date"]

    assert result.dt.strftime("%Y-%m-%d").tolist()[:4] == [
        "2025-01-15",
        "2025-02-10",
        "2025-03-05",
        "2025-04-07",
    ]
    assert pd.isna(result[4])


def test_clean_products_uses_most_common_spelling() -> None:
    df = text_frame(
        product=["Laptop Stand", "laptop  stand", " Laptop Stand ", "USB-C Hub"]
    )

    result = clean_products(df)["product"].tolist()

    assert result == ["Laptop Stand", "Laptop Stand", "Laptop Stand", "USB-C Hub"]


def test_remove_duplicates(caplog: pytest.LogCaptureFixture) -> None:
    df = text_frame(order_id=["1", "1", "2"], region=["West", "West", "East"])

    with caplog.at_level(logging.INFO):
        result = remove_duplicates(df)

    assert result["order_id"].tolist() == ["1", "2"]
    assert "Removed 1 duplicate rows. 2 rows left." in caplog.text


def test_fill_blanks() -> None:
    df = pd.DataFrame(
        {
            "quantity": ["2", "2", None, None],
            "unit_price": [10.0, None, 10.0, None],
            "revenue": [None, 30.0, 40.0, 50.0],
            "sales_rep": ["Ana Lee", None, None, None],
        }
    )

    result = fill_blanks(df)

    assert result["revenue"].tolist()[:3] == [20.0, 30.0, 40.0]
    assert result["unit_price"].tolist()[:3] == [10.0, 15.0, 10.0]
    assert result["quantity"].tolist()[:3] == [2, 2, 4]
    # Only one of the three numbers is known, so nothing can be worked out
    assert pd.isna(result.loc[3, "quantity"])
    assert pd.isna(result.loc[3, "unit_price"])
    assert result["sales_rep"].tolist() == ["Ana Lee", *["Unassigned"] * 3]


def test_find_missing() -> None:
    df = pd.DataFrame(
        {
            "order_id": ["1", "2", "3"],
            "order_date": pd.to_datetime(["2025-01-15", "2025-01-16", None]),
            "region": ["West", "East", "West"],
            "product": ["Desk Lamp", "Desk Lamp", "USB-C Hub"],
            "quantity": [2.0, None, 1.0],
            "unit_price": [10.0, None, 5.0],
            "revenue": [20.0, 30.0, 5.0],
            "sales_rep": ["Ana Lee", "Ben Cho", "Ana Lee"],
        }
    )

    result = find_missing(df)

    assert result["order_id"].tolist() == ["2", "3"]
    assert result["missing"].tolist() == ["quantity, unit_price", "order_date"]


def test_find_missing_when_every_row_is_complete(tmp_path: Path) -> None:
    df = clean_data(load_data(write_csv(tmp_path)))

    result = find_missing(df)

    assert result.empty
    assert "missing" in result.columns


def test_build_summaries(tmp_path: Path) -> None:
    df = clean_data(load_data(write_csv(tmp_path)))

    summaries = build_summaries(df)

    assert summaries.by_month["revenue"].tolist() == [90.0, 45.0, 79.0]
    assert summaries.by_month["change_pct"].tolist()[1:] == [-50.0, 75.6]
    assert summaries.by_product["revenue"].to_dict() == {
        "Laptop Stand": 135.0,
        "USB-C Hub": 79.0,
    }
    assert summaries.by_region["revenue"].to_dict() == {
        "West": 90.0,
        "East": 79.0,
        "South": 45.0,
    }


def test_build_report_writes_file(tmp_path: Path) -> None:
    output = tmp_path / "reports" / "sales_report.xlsx"

    df = build_report(write_csv(tmp_path), output)

    assert len(df) == 3
    workbook = load_workbook(output)
    assert workbook.sheetnames == [
        "Monthly",
        "Products",
        "Regions",
        "Clean Data",
        "Needs Review",
    ]
    assert workbook["Monthly"]["B1"].value == "Revenue"
    assert workbook["Products"]["A2"].value == "Laptop Stand"


def test_build_report_writes_to_memory(tmp_path: Path) -> None:
    output = BytesIO()

    build_report(write_csv(tmp_path), output)

    workbook = load_workbook(BytesIO(output.getvalue()))
    assert workbook["Clean Data"].max_row == 4


def test_build_report_on_sample_file(tmp_path: Path) -> None:
    df = build_report(SAMPLE_FILE, tmp_path / "sales_report.xlsx")

    # The numbers in the README's "Results on the sample file" table
    assert len(df) == 600
    assert df["region"].nunique() == 4
    assert df["product"].nunique() == 8
    assert df["quantity"].isna().sum() == 1
    assert df["unit_price"].isna().sum() == 4
    assert df["revenue"].isna().sum() == 5
    assert df["sales_rep"].isna().sum() == 0
    assert len(find_missing(df)) == 5


def test_report_lists_rows_needing_review(tmp_path: Path) -> None:
    output = tmp_path / "sales_report.xlsx"

    build_report(SAMPLE_FILE, output)

    rows = list(load_workbook(output)["Needs Review"].values)
    assert rows[0][-1] == "missing"
    assert len(rows) == 6
    assert rows[1][0] == "ORD-1469"
    assert rows[1][-1] == "unit_price, revenue"


def test_data_sheets_show_dollars_and_fit_long_names(tmp_path: Path) -> None:
    output = tmp_path / "sales_report.xlsx"

    build_report(SAMPLE_FILE, output)

    workbook = load_workbook(output)
    clean = workbook["Clean Data"]
    review = workbook["Needs Review"]
    # Unit price and revenue are columns F and G
    assert clean["F2"].number_format == "$#,##0.00"
    assert clean["G2"].number_format == "$#,##0.00"
    assert review["F6"].value == 32.75
    assert review["F6"].number_format == "$#,##0.00"
    # Product is column D, wide enough for "Noise-Canceling Headphones"
    assert clean.column_dimensions["D"].width >= 28
    assert review.column_dimensions["D"].width >= 28


def test_main_saves_report(tmp_path: Path, caplog: pytest.LogCaptureFixture) -> None:
    output = tmp_path / "sales_report.xlsx"

    with caplog.at_level(logging.INFO):
        exit_code = main([str(write_csv(tmp_path)), "--output", str(output)])

    assert exit_code == 0
    assert output.exists()
    assert f"Report saved to {output}" in caplog.text


def test_main_reports_errors(tmp_path: Path, caplog: pytest.LogCaptureFixture) -> None:
    exit_code = main([str(tmp_path / "missing.csv")])

    assert exit_code == 1
    assert "File not found" in caplog.text


def test_runs_as_script(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    script = Path(__file__).parent.parent / "generate_report.py"
    output = tmp_path / "sales_report.xlsx"
    monkeypatch.setattr(
        sys, "argv", ["generate_report.py", str(write_csv(tmp_path)), "-o", str(output)]
    )

    with pytest.raises(SystemExit) as exit_info:
        runpy.run_path(str(script), run_name="__main__")

    assert exit_info.value.code == 0
    assert output.exists()
